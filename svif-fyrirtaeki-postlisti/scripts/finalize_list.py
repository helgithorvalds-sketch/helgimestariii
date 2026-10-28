#!/usr/bin/env python3
"""Turns the crawler output (plus the registry fallbacks) into the mailing list: one row per company with the
address that will be written to, who it is, how it was verified, the greeting and the category phrase that go
into the e-mail, and empty columns for the sending log.

Order of preference for the address:
  1. markaðsstjóri found on the company's site      4. sölustjóri found on the site
  2. framkvæmdastjóri / eigandi found on the site   5. a named person on the site (title unknown)
  3. a named lawyer / auditor from the registry     6. the company's general address found on the site

Usage: finalize_list.py contacts.csv postlisti.csv
"""
import csv, re, sys, os, json
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from greeting import greeting, gender
import requests

SRC, DST = sys.argv[1], sys.argv[2]
clean = lambda s: re.sub(r'\s+', ' ', s or '').strip()

# the phrase that fills "...einmitt á þeim tímapunkti sem þau eru að velja sér {þjónusta} og annað sem nýr rekstur þarf"
PHRASES = [
    (r'lögfræ|lögm|hugverk|vörumerkjaskr|einkaleyf', 'lögfræðiþjónustu'),
    (r'endurskoð', 'endurskoðanda'),
    (r'bókhald|bókar', 'bókhaldsstofu'),
    (r'launa|rekstrarhugbún|bókhaldskerf|hugbúnað', 'launa- og bókhaldskerfi'),
    (r'innheimt', 'innheimtuþjónustu'),
    (r'bank|greiðslu|pos', 'banka og greiðslulausnir'),
    (r'trygg|vátrygg', 'tryggingar'),
    (r'fjármögn|rekstrarleig|leigufélag', 'fjármögnun'),
    (r'auglýsingastof|birting', 'auglýsingastofu'),
    (r'hönnun|grafísk|vörumerk|ásýnd', 'hönnuð fyrir merki og ásýnd'),
    (r'vefstof|vefsíð|vefhönn|vefversl', 'vefsíðu'),
    (r'stafræn|seo|samfélagsmið|markaðssetn', 'markaðsþjónustu'),
    (r'ljósmynd|myndband|efnisgerð|drón', 'ljósmyndara og efnisgerð'),
    (r'prentsmið|prentþjón|prentun', 'prentþjónustu'),
    (r'merking|skilt|kynningarvör|fatnað', 'merkingar og skilti'),
    (r'tölvubún|tölvuversl', 'tölvubúnað'),
    (r'upplýsingatæk|rekstrarþjón|tölvuþjón|hýsing', 'tölvuþjónustu'),
    (r'sölukerf|netverslunarkerf', 'sölu- og bókhaldskerfi'),
    (r'fjarskipt|símkerf', 'síma og fjarskipti'),
    (r'prentar|afritun|skrifstofutæk', 'skrifstofutæki'),
    (r'netörygg|persónuvernd|undirrit', 'netöryggi og rafræna undirritun'),
    (r'lén|tölvupóst', 'lén og vefhýsingu'),
    (r'skrifstofuhúsn|samvinnurým|skrifstofuhót|atvinnuhúsn|fasteign', 'skrifstofuhúsnæði'),
    (r'húsgögn|innrétting', 'skrifstofuhúsgögn'),
    (r'rekstrarvör|pappír|skrifstofuvör', 'rekstrarvörur'),
    (r'kaffi', 'kaffivélar'),
    (r'ræsting|þrif', 'ræstingar'),
    (r'örygg|brunavarn|slökkvi', 'öryggiskerfi'),
    (r'bíl|sendibíl|flutning', 'bíla og flutninga'),
    (r'mötuneyt|matar|veiting', 'matarþjónustu'),
    (r'ráðning|mannauð', 'ráðningarþjónustu'),
    (r'markaðsráðgj|söluþjálf|rannsókn', 'markaðsráðgjöf'),
    (r'þýðing|texta', 'þýðingar'),
    (r'námskeið|fræðsl|þjálfun', 'fræðslu og námskeið'),
    (r'vinnuvernd|heilsa|heilsurækt|líkamsrækt', 'heilsu- og vinnuverndarþjónustu'),
    (r'stofnun|styrkj|viðskiptaáætl', 'ráðgjöf við stofnun'),
    (r'ráðgjöf|stjórnend', 'ráðgjöf'),
]
def phrase(flokkur, undir):
    t = (flokkur + ' ' + undir).lower()
    for rx, ph in PHRASES:
        if re.search(rx, t): return ph
    return 'þjónustu eins og ykkar'

NOUNISH = re.compile(r'^(Bókhald|Endurskoðun|Ráðgjöf|Lögfræði|Skrifstofa|Starfsfólk|Stjórn|Teymi|Hönnun|Prentun|Verslun|Þjónusta|Netfang|Sími|Nafn|Reykjavík|Akureyri|Ísland|Sala|Markaður|Markaðs|Fyrirtæki|Hafðu|Hafa|Samband)$', re.I)
MX = {}
def mx(domain):
    if domain in MX: return MX[domain]
    try:
        j = requests.get('https://dns.google/resolve', params={'name': domain, 'type': 'MX'}, timeout=15).json()
        ans = [a['data'] for a in j.get('Answer', []) if a.get('type') == 15]
        MX[domain] = 'MX:' + ans[0].split()[-1].rstrip('.') if ans else 'ekkert MX'
    except Exception:
        MX[domain] = 'DNS?'
    return MX[domain]

NOUN_WORDS = re.compile(r'^(Bókhald|Endurskoðun|Ráðgjöf|Lögfræði|Lögmenn|Lögmaður|Skrifstofa|Starfsfólk|Starfsmaður|Yfirstjórn|Fjármáladeild|Söludeild|Markaðsdeild|Deild|Stjórn|Teymi|Hönnun|Prentun|Verslun|Þjónusta|Símatími|Opið|Hafðu|Hafa|Samband|Netfang|Sími|Símanúmer|Skilaboð|Senda|Nafn|Fax|Kennitala|Heimilisfang|Fyrirtæki|Reykjavík|Kópavogur|Hafnarfjörður|Akureyri|Garðabær|Mosfellsbær|Selfoss|Ísland|Iceland|Google|Facebook|Instagram|Linkedin|About|Services|Service|Contact|Team|Home|Menu|Email|Phone|Partner|Partners|Ads|Meta|Microsoft|Apple|Premium|Vefur|Vefsíða|Vefverslun|Tölvupóstur|Póstur|Sala|Markaður|Markaðs)$', re.I)
TITLE_WORD = re.compile(r'(stjóri|stjórnandi|eigandi|stofnandi|fulltrúi|ráðgjafi|endurskoðandi|lögmaður|lögfræðingur|bókari|hönnuður|sérfræðingur|forritari|þjónn|CEO|CFO|COO|CMO|Partner|Manager|Managing|Director|Senior|Head|Chief|Lead|Founder|Owner|Consultant|Designer|Developer|Specialist)$', re.I)
def person_name(n):
    """Strip labels and title words that the crawler left around a person's name; '' when nothing person-like remains."""
    w = [x for x in clean(n).split() if x]
    while w and (NOUN_WORDS.match(w[0]) or TITLE_WORD.search(w[0])): w = w[1:]
    while w and (NOUN_WORDS.match(w[-1]) or TITLE_WORD.search(w[-1])): w = w[:-1]
    w = w[:4]
    if not w or len(w[0]) < 3 or not re.match(r'^[A-ZÁÉÍÓÚÝÞÆÖ][a-záéíóúýþæöð]+$', w[0]): return ''
    return ' '.join(w)

def company_display(name):
    n = clean(name)
    n = re.sub(r'\s*[-–]\s*(Reykjavík|Akureyri)$', '', n)
    n = re.sub(r'\s+(ehf|hf|slf|sf|ohf|svf)\.?$', '', n, flags=re.I)
    return n

out = []
for r in csv.DictReader(open(SRC, encoding='utf-8'), delimiter=';'):
    typ = r.get('tegund', ''); site_name = r.get('tengilidur', '')
    reg_name, reg_mail, reg_title = r.get('skra_tengilidur', ''), r.get('skra_netfang', ''), r.get('skra_titill', '')
    choice = None
    if r.get('netfang') and typ in ('markadur', 'stjornandi'):
        choice = (r['netfang'], site_name, r.get('titill', ''), {'markadur': 'markaðsstjóri', 'stjornandi': 'stjórnandi/eigandi'}[typ], 'á vef fyrirtækisins: ' + r.get('netfang_heimild', ''))
    elif reg_mail:
        choice = (reg_mail, reg_name, reg_title, 'nafngreindur ' + reg_title, 'í félagatali: ' + r.get('heimild', '').split(' | ')[0])
    elif r.get('netfang') and typ in ('sala', 'nafngreint'):
        choice = (r['netfang'], site_name, r.get('titill', ''), 'sölustjóri' if typ == 'sala' else 'nafngreindur starfsmaður', 'á vef fyrirtækisins: ' + r.get('netfang_heimild', ''))
    elif r.get('almennt_netfang'):
        choice = (r['almennt_netfang'], '', '', 'almennt netfang', 'á vef fyrirtækisins: ' + r.get('netfang_heimild', ''))
    elif r.get('netfang'):
        choice = (r['netfang'], site_name, r.get('titill', ''), 'netfang á vef (óflokkað)', 'á vef fyrirtækisins: ' + r.get('netfang_heimild', ''))
    row = dict(fyrirtaeki=company_display(r['fyrirtaeki']), flokkur=r.get('flokkur', ''), undirflokkur=r.get('undirflokkur', ''), vefur=r.get('vefur_endanlegur') or r.get('vefur', ''),
               netfang='', tengilidur='', titill='', tegund='', stadfesting='', mx='', kvedja='', thjonusta=phrase(r.get('flokkur', ''), r.get('undirflokkur', '')),
               onnur_netfong=r.get('oll_netfong', '')[:300], athugasemd=r.get('villa', '') or r.get('athugasemd', ''), sent='', message_id='', svar='')
    if choice:
        e, n, t, typ_is, verif = choice
        n = person_name(n)
        row.update(netfang=e, tengilidur=n, titill=t, tegund=typ_is, stadfesting=verif, mx=mx(e.split('@')[1]), kvedja=greeting(n))
    else:
        row['athugasemd'] = (row['athugasemd'] + '; ' if row['athugasemd'] else '') + 'ekkert netfang fannst'
    out.append(row)

out.sort(key=lambda r: (r['netfang'] == '', r['flokkur'], r['fyrirtaeki']))
seen = {}
for r in out:
    e = r['netfang'].lower()
    if not e: continue
    if e in seen:
        r['athugasemd'] = (r['athugasemd'] + '; ' if r['athugasemd'] else '') + f"sama netfang og {seen[e]}, ekki sent tvisvar"
        r['sent'] = 'sleppt: sama netfang og ' + seen[e]
    else:
        seen[e] = r['fyrirtaeki']
with open(DST, 'w', newline='', encoding='utf-8') as f:
    w = csv.DictWriter(f, fieldnames=list(out[0].keys()), delimiter=';'); w.writeheader(); w.writerows(out)
# the same list as a spreadsheet, for reading and ticking off by hand
try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook(); ws = wb.active; ws.title = 'Póstlisti'
    cols = list(out[0].keys()); ws.append(cols)
    for c in ws[1]: c.font = Font(bold=True); c.fill = PatternFill('solid', fgColor='E8ECF5')
    for r in out: ws.append([r[c] for c in cols])
    widths = dict(fyrirtaeki=34, flokkur=28, undirflokkur=24, vefur=26, netfang=32, tengilidur=28, titill=22, tegund=24, stadfesting=50, mx=18, kvedja=16, thjonusta=24, onnur_netfong=40, athugasemd=30, sent=12, message_id=18, svar=14)
    for i, c in enumerate(cols, 1): ws.column_dimensions[get_column_letter(i)].width = widths.get(c, 16)
    ws.freeze_panes = 'B2'; ws.auto_filter.ref = ws.dimensions
    wb.save(DST.rsplit('.', 1)[0] + '.xlsx')
except Exception as e:
    print('xlsx skipped:', e)
n_mail = sum(1 for r in out if r['netfang'])
from collections import Counter
print(f'{len(out)} companies, {n_mail} with an address;', Counter(r['tegund'] for r in out if r['netfang']).most_common())
print('MX problems:', sum(1 for r in out if r['netfang'] and not r['mx'].startswith('MX:')))
